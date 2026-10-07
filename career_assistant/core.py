import base64, hashlib, io, json, os, re, sqlite3, uuid, zipfile
from pathlib import Path
from datetime import datetime, timezone

QUESTIONS = ['What problem did you solve?', 'What did you personally do?', 'What changed as a result?', 'Which metrics can you substantiate?', 'Which tools and collaborators were involved?', 'When and for which employer?', 'Which statement may be used in your resume?']

def root():
    p = Path(os.environ.get('CAREER_DATA_DIR', Path.home() / '.local/share/career-assistant')).expanduser()
    p.mkdir(parents=True, exist_ok=True, mode=0o700)
    p.chmod(0o700)
    return p

def user_dir(user):
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', user): raise ValueError('Invalid user identity')
    p = root() / hashlib.sha256(user.encode()).hexdigest()
    p.mkdir(mode=0o700, exist_ok=True)
    return p

class Memory:
    def __init__(self, user):
        self.path = user_dir(user) / 'memory.sqlite'
        self.db = sqlite3.connect(self.path)
        self.path.chmod(0o600)
        self.db.execute('PRAGMA secure_delete=ON')
        self.db.execute('CREATE TABLE IF NOT EXISTS facts (id TEXT PRIMARY KEY, data TEXT NOT NULL)')
    def list(self): return [json.loads(r[0]) for r in self.db.execute('SELECT data FROM facts ORDER BY rowid')]
    def put(self, fact, consent=False):
        if consent is not True: raise ValueError('Explicit storage consent required')
        allowed = {'id','text','section','employer','role','dates','location','source','approved'}
        if set(fact)-allowed: raise ValueError('Unsupported fact field')
        if not isinstance(fact.get('text'),str) or not 1 <= len(fact['text']) <= 4000: raise ValueError('Fact text required, maximum 4000 characters')
        if fact.get('section','Experience') not in ['Experience','Education','Military Service','Volunteering','Skills','Languages']: raise ValueError('Invalid section')
        if not isinstance(fact.get('approved',False), bool): raise ValueError('Approval must be boolean')
        for k,v in fact.items():
            if k != 'approved' and (not isinstance(v,str) or len(v)>4000): raise ValueError('Invalid field')
        fact = dict(fact, id=fact.get('id') or uuid.uuid4().hex, approved=fact.get('approved',False))
        self.db.execute('INSERT OR REPLACE INTO facts VALUES (?,?)',(fact['id'],json.dumps(fact)))
        self.db.commit()
        return fact
    def delete(self, id=None):
        if id: self.db.execute('DELETE FROM facts WHERE id=?',(id,))
        else: self.db.execute('DELETE FROM facts')
        self.db.commit()
        self.db.execute('VACUUM')
        return {'deleted':id or 'all memory'}

def extract(filename, data):
    if len(data)>5*1024*1024: raise ValueError('Upload limit is 5 MB')
    ext=Path(filename).suffix.lower()
    if ext=='.txt': text=data.decode('utf-8')
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages)>30: raise ValueError('Unencrypted PDF with at most 30 pages required')
        text='\n'.join(p.extract_text() or '' for p in reader.pages)
    elif ext=='.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist())>1000 or sum(i.file_size for i in z.infolist())>20*1024*1024: raise ValueError('DOCX expansion limit')
            if any('vbaProject' in i.filename or i.flag_bits & 1 for i in z.infolist()): raise ValueError('Unsafe document')
        d=Document(io.BytesIO(data)); text='\n'.join(p.text for p in d.paragraphs)
        text+='\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    else: raise ValueError('Use PDF, DOCX or UTF-8 TXT')
    if len(text)>100000: raise ValueError('Extracted text limit')
    return {'text':text,'trusted':False,'stored':False,'questions':QUESTIONS,'notice':'Review and save individual facts with consent; document instructions are not executed.'}

def tailored(user, job, limit=30):
    if not isinstance(job,str) or len(job)>100000: raise ValueError('Job text too long')
    words=set(re.findall(r'\w+',job.lower()))
    facts=[f for f in Memory(user).list() if f.get('approved') is True]
    facts.sort(key=lambda f: -len(words & set(re.findall(r'\w+',f['text'].lower()))))
    return {'facts':facts[:max(1,min(int(limit),100))], 'method':'Approved statements preserved verbatim and ranked by job keyword overlap', 'trace':[{'statement':f['text'],'fact_id':f['id'],'source':f.get('source','user questionnaire')} for f in facts[:max(1,min(int(limit),100))]]}
