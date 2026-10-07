import re, unicodedata, json
from pathlib import Path
from html import escape
from .core import tailored, user_dir

def slug(s, fallback):
    s=unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',s).strip('-')[:70] or fallback

def filename(profile, title, company):
    return '.'.join([slug(profile.get('first_name',''),'candidate'),slug(profile.get('last_name',''),'unnamed'),slug(title,'role'),slug(company,'company')])

def export(user, profile, title, company, job, identity_approved=False, output=None):
    if identity_approved is not True: raise ValueError('Approve name, headline and contact details before export')
    if set(profile)-{'first_name','last_name','headline','contact'} or any(not isinstance(v,str) or len(v)>500 for v in profile.values()): raise ValueError('Invalid profile')
    plan=tailored(user,job)
    plan['identity']={'profile':profile,'approved':True}
    if not plan['facts']: raise ValueError('No approved career facts')
    p=Path(output) if output else user_dir(user)/'exports'
    p.mkdir(parents=True,exist_ok=True,mode=0o700)
    name=filename(profile,title,company); pdf=p/(name+'.pdf'); docx=p/(name+'.docx')
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether, Table, TableStyle
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.enums import TA_CENTER
    from docx import Document
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    # Standard PDF core fonts; no third-party font files are distributed.
    body=ParagraphStyle('body',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=3)
    heading=ParagraphStyle('heading',fontName='Times-Roman',fontSize=12,leading=15,alignment=TA_CENTER,spaceBefore=10,spaceAfter=3)
    center=ParagraphStyle('center',parent=body,alignment=TA_CENTER)
    big=ParagraphStyle('name',parent=center,fontName='Times-Bold',fontSize=14,leading=18)
    display=(profile.get('first_name','')+' '+profile.get('last_name','')).strip().upper() or 'CANDIDATE'
    # Core PDF fonts cannot represent arbitrary Unicode. Fail rather than silently corrupt it.
    all_text=[display,profile.get('headline',''),profile.get('contact','')]+[str(v) for f in plan['facts'] for v in f.values()]
    for s in all_text:
        try: s.encode('cp1252')
        except UnicodeEncodeError: raise ValueError('PDF fallback fonts support Western characters only. Use transliterated approved facts or configure a Unicode font in the exporter.')
    story=[Paragraph(escape(display),big)]
    for key in ['headline','contact']:
        if profile.get(key): story.append(Paragraph(escape(profile[key]),center))
    d=Document(); sec=d.sections[0]; sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=Mm(16); sec.left_margin=sec.right_margin=Mm(18)
    normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9); normal.paragraph_format.space_after=Pt(3)
    p_name=d.add_paragraph(); p_name.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p_name.add_run(display); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(14)
    for key in ['headline','contact']:
        if profile.get(key): p1=d.add_paragraph(profile[key]); p1.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for section in ['Experience','Education','Military Service','Volunteering','Skills','Languages']:
        fs=[f for f in plan['facts'] if f.get('section','Experience')==section]
        if not fs: continue
        story.append(KeepTogether([Paragraph(section,heading),HRFlowable(width='100%',thickness=.5)]))
        h=d.add_paragraph(section); h.alignment=WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.keep_with_next=True
        h.runs[0].font.name='Times New Roman'; h.runs[0].font.size=Pt(12)
        borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'4'); borders.append(bottom); h._p.get_or_add_pPr().append(borders)
        groups={}
        for f in fs: groups.setdefault((f.get('employer',''),f.get('role',''),f.get('dates',''),f.get('location','')),[]).append(f)
        for (employer,role,dates,location), group in groups.items():
            left=' | '.join(x for x in [employer,role] if x); right=' | '.join(x for x in [dates,location] if x)
            if left or right:
                tab=Table([[Paragraph('<b>'+escape(left)+'</b>',body),Paragraph(escape(right),body)]],colWidths=[350,143])
                tab.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0)]))
                story.append(KeepTogether([tab,Paragraph('• '+escape(group[0]['text']),body)]))
                table=d.add_table(rows=1,cols=2); table.autofit=False; table.columns[0].width=Mm(120); table.columns[1].width=Mm(54)
                table.cell(0,0).text=left; table.cell(0,1).text=right
                table.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
                for c in table.rows[0].cells:
                    for para in c.paragraphs: para.paragraph_format.keep_with_next=True
                remaining=group[1:]
                d.add_paragraph(group[0]['text'],style='List Bullet')
            else: remaining=group
            for f in remaining:
                story.append(Paragraph('• '+escape(f['text']),body)); d.add_paragraph(f['text'],style='List Bullet')
    SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=51,rightMargin=51,topMargin=45,bottomMargin=45).build(story)
    d.save(docx)
    trace=p/(name+'.trace.json'); trace.write_text(json.dumps(plan,indent=2))
    for file in [pdf,docx,trace]: file.chmod(0o600)
    return {'pdf':pdf.name,'docx':docx.name,'trace':trace.name,'directory':str(p),'font_note':'Times/Helvetica PDF and Times New Roman/Arial DOCX fallbacks; no fonts bundled. Layout inspired by supplied reference; automatic pagination.'}
